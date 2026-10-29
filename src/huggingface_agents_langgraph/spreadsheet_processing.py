"""Spreadsheet processing node for tasks with Excel attachments."""

import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, TypedDict

from openpyxl import load_workbook

from huggingface_agents_langgraph.llm import ChatMessage, chat_completion
from huggingface_agents_langgraph.state import AgentState, Evidence


class SpreadsheetProcessingUpdate(TypedDict):
    """State fields produced by the spreadsheet processing node."""

    evidence: list[Evidence]


def spreadsheet_processing(state: AgentState) -> SpreadsheetProcessingUpdate:
    """Use deterministic calculations when possible and retain LLM fallback."""
    file_path = state.get("file_path")
    if file_path is None:
        raise ValueError("file_path is required to process an Excel attachment")

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel attachment does not exist: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {
            worksheet.title: [
                [cell for cell in row]
                for row in worksheet.iter_rows(values_only=True)
                if any(cell is not None for cell in row)
            ]
            for worksheet in workbook.worksheets
        }
    finally:
        workbook.close()

    if not any(sheets.values()):
        raise ValueError("Excel attachment contains no non-empty cells")

    workbook_data = json.dumps(sheets, ensure_ascii=False, default=str)
    plan_messages: list[ChatMessage] = [
        {
            "role": "system",
            "content": (
                "Decide whether the spreadsheet question can be expressed as one "
                "simple deterministic calculation. If so, return only a JSON "
                "object with strategy='compute', sheet, operation (sum, count, "
                "average, min, or max), column (header text, or null for count), "
                "and filters (a list of {column, operator, value}; operators are "
                "equals, not_equals, contains, greater_than, greater_or_equal, "
                "less_than, less_or_equal). Assume the first non-empty row is the "
                "header. If the task needs a more complex interpretation, "
                "comparison, ranking, multi-step calculation, or anything this "
                "simple plan cannot represent, return only "
                '{"strategy":"analyze"}. Never simplify a question to fit '
                "the calculation plan. /no_think"
            ),
        },
        {
            "role": "user",
            "content": (
                f"Question:\n{state['question']}\n\nWorkbook sheets and rows:\n"
                f"{workbook_data}"
            ),
        },
    ]
    plan_response = chat_completion(plan_messages, max_tokens=384)
    try:
        plan = json.loads(plan_response)
    except json.JSONDecodeError:
        plan = {"strategy": "analyze"}

    if isinstance(plan, dict) and plan.get("strategy") == "compute":
        result = _compute_plan(sheets, plan)
        if result is not None:
            evidence: Evidence = {
                "content": result,
                "source": str(path),
                "source_type": "spreadsheet_calculation",
            }
            return {"evidence": [evidence]}

    return _analyze_workbook(state["question"], workbook_data, path)


def _compute_plan(
    sheets: dict[str, list[list[Any]]], plan: dict[str, Any]
) -> str | None:
    """Run a supported aggregation plan; return None for unsupported plans."""
    sheet_name = plan.get("sheet")
    operation = plan.get("operation")
    column = plan.get("column")
    filters = plan.get("filters", [])
    if (
        sheet_name not in sheets
        or operation not in {"sum", "count", "average", "min", "max"}
        or not isinstance(filters, list)
        or (operation != "count" and not isinstance(column, str))
    ):
        return None

    rows = sheets[sheet_name]
    if len(rows) < 2 or not rows[0]:
        return None
    headers = [str(header).strip() if header is not None else "" for header in rows[0]]
    if len(set(headers)) != len(headers):
        return None
    column_indexes = {header: index for index, header in enumerate(headers) if header}
    if operation != "count" and column not in column_indexes:
        return None

    filter_indexes: list[tuple[int, str, Any]] = []
    for condition in filters:
        if not isinstance(condition, dict):
            return None
        filter_column = condition.get("column")
        operator = condition.get("operator")
        if (
            filter_column not in column_indexes
            or operator
            not in {
                "equals",
                "not_equals",
                "contains",
                "greater_than",
                "greater_or_equal",
                "less_than",
                "less_or_equal",
            }
            or "value" not in condition
        ):
            return None
        filter_indexes.append(
            (column_indexes[filter_column], operator, condition["value"])
        )

    matching_rows = [
        row
        for row in rows[1:]
        if all(
            _matches(row[index] if index < len(row) else None, operator, value)
            for index, operator, value in filter_indexes
        )
    ]
    if operation == "count":
        return f"Count: {len(matching_rows)}"

    value_index = column_indexes[column]
    try:
        values = [
            Decimal(str(row[value_index]))
            for row in matching_rows
            if value_index < len(row) and row[value_index] is not None
        ]
    except InvalidOperation:
        return None
    if not values:
        return None

    if operation == "sum":
        result = sum(values, Decimal(0))
    elif operation == "average":
        result = sum(values, Decimal(0)) / len(values)
    elif operation == "min":
        result = min(values)
    else:
        result = max(values)
    return f"{operation.title()} of {column}: {result}"


def _matches(cell: Any, operator: str, expected: Any) -> bool:
    """Compare a cell with one supported filter condition."""
    if operator == "equals":
        return str(cell).casefold() == str(expected).casefold()
    if operator == "not_equals":
        return str(cell).casefold() != str(expected).casefold()
    if operator == "contains":
        return str(expected).casefold() in str(cell).casefold()
    try:
        actual_number = Decimal(str(cell))
        expected_number = Decimal(str(expected))
    except InvalidOperation:
        return False
    if operator == "greater_than":
        return actual_number > expected_number
    if operator == "greater_or_equal":
        return actual_number >= expected_number
    if operator == "less_than":
        return actual_number < expected_number
    return actual_number <= expected_number


def _analyze_workbook(
    question: str, workbook_data: str, path: Path
) -> SpreadsheetProcessingUpdate:
    """Use the model's general analysis ability for plans beyond the DSL."""
    messages: list[ChatMessage] = [
        {
            "role": "system",
            "content": (
                "Answer the spreadsheet question using the workbook data. This "
                "fallback supports complex interpretation and calculations that "
                "do not fit a simple aggregation plan. Carefully show the relevant "
                "findings and calculation result, but do not invent missing data. "
                "If the workbook does not contain enough information, say so."
            ),
        },
        {
            "role": "user",
            "content": (
                f"Question:\n{question}\n\nWorkbook sheets and rows:\n{workbook_data}"
            ),
        },
    ]
    evidence: Evidence = {
        "content": chat_completion(messages),
        "source": str(path),
        "source_type": "spreadsheet_analysis",
    }
    return {"evidence": [evidence]}
